"""Complete zero-inclusive source-grain/field profile for UK PoMS FIT Count.

Never reconstruct negative observations for studies from positive-only network data:
FIT count is a separate prospective focal-flower observation design.
"""
from __future__ import annotations
import io
import json
from pathlib import Path
import re
import zipfile
import numpy as np
import pandas as pd
from stage3_source_audit import DATASET_DOI

TABLES = {
    'public': 'data/ukpoms_publicfitcountdata_2017-2023.csv',
    'one_km': 'data/ukpoms_1kmfitcountdata_2017-2023.csv'
}
GROUPS = ['bumblebees','honeybees','solitary_bees','wasps','hoverflies',
          'other_flies','butterflies_moths','beetles','insects_small','insects_other']
COUNT_COLS = GROUPS+['all_insects_total','floral_unit_count']

def profile_table(df, surface):
    if not {'sample_id','date','year','target_flower','floral_unit_count',
            'wind_speed','all_insects_total',*GROUPS} <= set(df.columns):
        raise ValueError('mandatory FIT survey and insect columns missing')
    n=len(df)
    num = df[COUNT_COLS].apply(pd.to_numeric,errors='coerce')
    invalid={c:int((num[c].notna() & ((num[c]<0) | (num[c]%1!=0))).sum()) for c in COUNT_COLS}
    if any(invalid.values()):raise ValueError(f'negative/fractional counts not valid: {invalid}')
    complete=num[GROUPS+['all_insects_total']].notna().all(axis=1)
    consistency=num.loc[complete,GROUPS].sum(axis=1)-num.loc[complete,'all_insects_total']
    date=pd.to_datetime(df['date'],errors='coerce',dayfirst=True)
    d={
        'table':surface,'rows':n,
        'sample_ids_unique':int(df.sample_id.nunique(dropna=True)),
        'sample_id_duplicated_rows':int(df.sample_id.duplicated(keep=False).sum()),
        'years':{str(k):int(v) for k,v in df.year.value_counts(dropna=False).sort_index().items()},
        'bad_date':int(date.isna().sum()),
        'min_date':str(date.min()),'max_date':str(date.max()),
        'zero_total_insect_rows':int(num.all_insects_total.eq(0).sum()),
        'zero_total_rate':float(num.all_insects_total.eq(0).mean()),
        'positive_total_rows':int(num.all_insects_total.gt(0).sum()),
        'total_insects_count':float(num.all_insects_total.sum()),
        'total_consistency_complete_rows':int(complete.sum()),
        'total_insect_group_sum_mismatches':int((consistency!=0).sum()),
        'insect_total_mismatch_example_diffs':sorted(consistency.loc[consistency!=0].unique().tolist())[:8],
        'recorded_flower_units_zero':int(num.floral_unit_count.eq(0).sum()),
        'recorded_flower_units_missing':int(num.floral_unit_count.isna().sum()),
        'flower_units_median_positive':float(num.loc[num.floral_unit_count.gt(0),'floral_unit_count'].median()),
        'flower_cover_levels':{str(k):int(v) for k,v in df.flower_cover.value_counts(dropna=False).head(20).items()},
        'wind_speed_levels':{str(k):int(v) for k,v in df.wind_speed.value_counts(dropna=False).head(30).items()},
        'sunshine_levels':{str(k):int(v) for k,v in df.sunshine.value_counts(dropna=False).head(30).items()},
        'cloud_cover_levels':{str(k):int(v) for k,v in df.cloud_cover.value_counts(dropna=False).head(30).items()},
        'target_flowers_top':{str(k):int(v) for k,v in df.target_flower_corrected.value_counts(dropna=False).head(24).items()},
        'target_flower_distinct':int(df.target_flower_corrected.nunique(dropna=True)),
        'floral_unit_levels':{str(k):int(v) for k,v in df.floral_unit.value_counts(dropna=False).head(20).items()},
        'country_levels':{str(k):int(v) for k,v in df.country.value_counts(dropna=False).head(20).items()},
        'recorder_types':{str(k):int(v) for k,v in df.recorder_type.value_counts(dropna=False).head(20).items()},
        'by_group':{c:{
            'sum':float(num[c].sum()),'zero_fraction':float(num[c].eq(0).mean()),
            'missing':int(num[c].isna().sum())} for c in GROUPS},
        'required_missing':{c:int(df[c].isna().sum()) for c in [
            'sample_id','date','year','target_flower_corrected','floral_unit_count',
            'wind_speed','sunshine','count_start_time','all_insects_total'
        ]}
    }
    for idc in ['x1km_square','location_code']:
        if idc in df.columns:
            d[idc+'_distinct']=int(df[idc].nunique(dropna=True))
            if idc=='location_code':
                times=df.groupby([idc,'year'],dropna=False).date.nunique()
                d['site_year_units']=int(len(times))
                d['site_year_with_ge3_dates']=int((times>=3).sum())
                d['site_year_with_ge4_dates']=int((times>=4).sum())
                x=df.groupby([idc,'year'],dropna=False).target_flower_corrected.nunique()
                d['site_year_with_ge2_target_flower_labels']=int((x>=2).sum())
    return d

def parse_docx_metadata(content):
    from docx import Document
    doc=Document(io.BytesIO(content))
    terms=['floral_unit','flower_cover','wind_speed','all_insects_total',
           'sample_id','target_flower','no insect','zero','wind speed','cloud_cover',
           'sunshine','digitised_by','target flower']
    lines=[]
    for para in doc.paragraphs:
        text=' '.join(para.text.split())
        if text and any(t in text.lower() for t in terms):lines.append(text[:1200])
    for table in doc.tables:
        for row in table.rows:
            text=' | '.join(' '.join(cell.text.split()) for cell in row.cells)
            if any(t in text.lower() for t in terms):lines.append(text[:1600])
    return {'relevant_metadata_excerpt':lines[:100], 'matched_blocks':len(lines),
            'document_paragraphs':len(doc.paragraphs),'document_tables':len(doc.tables)}

def read_original_csv_bytes(raw: bytes):
    """Strict decode, with an explicit non-lossy compatibility fallback."""
    for encoding in ('utf-8', 'cp1252', 'latin-1'):
        try:
            text = raw.decode(encoding, errors='strict')
            return pd.read_csv(io.StringIO(text), low_memory=False), encoding
        except UnicodeDecodeError:
            continue
    raise ValueError('no reversible text decoder available')

def main():
    import argparse
    ap=argparse.ArgumentParser()
    ap.add_argument('--out',default='stage3_poms_audit')
    args=ap.parse_args()
    out=Path(args.out)
    archive=out/'poms_fit_2017_2023_source.zip'
    if not archive.is_file():raise ValueError('official source download missing')
    report={'doi':DATASET_DOI,'status':'full_original_dataset_source_profile',
            'data_source':'Official original provider ZIP, not a positive-only visit extraction',
            'samples_are_target_flower_surveys_not_animal_decisions':True}
    with zipfile.ZipFile(archive) as z:
        ids={}
        for surface,member in TABLES.items():
            df, source_encoding = read_original_csv_bytes(z.read(member))
            ids[surface]=set(df.sample_id.dropna())
            report[surface]=profile_table(df,surface)
            report[surface]['source_text_encoding']=source_encoding
        report['sample_id_overlap_public_1km']=len(ids['public']&ids['one_km'])
        metadata='supporting-documents/ukpoms_fitcountdata_2017-2023_metadata.docx'
        report['field_metadata']=parse_docx_metadata(z.read(metadata))
    (out/'full_zero_inclusive_profile.json').write_text(json.dumps(report,indent=2,ensure_ascii=False))
    print('POMS_FULL_PROFILE '+json.dumps({k:v for k,v in report.items() if k!='field_metadata'}),flush=True)
    print('POMS_FIELD_METADATA '+json.dumps(report['field_metadata'],ensure_ascii=False),flush=True)
    print('POMS_FULL_PROFILE_COMPLETE',flush=True)

if __name__=='__main__':main()
