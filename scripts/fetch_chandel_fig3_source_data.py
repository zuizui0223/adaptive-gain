"""Fetch and inspect Chandel et al. (2024) Nature Source Data Fig. 3.

This is a one-off public-data Phase A helper. It finds the publisher/PMC
attachment link by anchor text, downloads the XLSX, then exports every worksheet
as CSV plus a machine-readable manifest and text preview.

The script intentionally does not infer individual mosquito latencies from
population-level IHSI source data.
"""
from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ARTICLE_URLS = (
    "https://www.nature.com/articles/s41586-024-07848-5",
    "https://pmc.ncbi.nlm.nih.gov/articles/PMC11410652/",
)

TARGET_PATTERNS = (
    re.compile(r"source\s*data\s*fig(?:ure)?\.?\s*3\b", re.I),
    re.compile(r"source\s*data.*fig(?:ure)?\.?\s*3\b", re.I),
)

UA = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "Chrome/130.0 Safari/537.36 adaptive-gain-public-data-audit/1.0"
)


def fetch(url: str) -> requests.Response:
    response = requests.get(
        url,
        headers={"User-Agent": UA, "Accept": "*/*"},
        timeout=60,
        allow_redirects=True,
    )
    response.raise_for_status()
    return response


def norm(text: str) -> str:
    return " ".join(text.replace("\xa0", " ").split())


def find_attachment() -> tuple[str, str]:
    diagnostics = []
    for page_url in ARTICLE_URLS:
        try:
            response = fetch(page_url)
        except Exception as exc:
            diagnostics.append({"page": page_url, "error": repr(exc)})
            continue

        soup = BeautifulSoup(response.text, "html.parser")
        candidates = []
        for anchor in soup.find_all("a", href=True):
            text = norm(anchor.get_text(" ", strip=True))
            title = norm(anchor.get("title", ""))
            aria = norm(anchor.get("aria-label", ""))
            searchable = " | ".join(x for x in (text, title, aria) if x)
            href = urljoin(response.url, anchor["href"])
            if any(pattern.search(searchable) for pattern in TARGET_PATTERNS):
                candidates.append((searchable, href))

        if not candidates:
            # Nature sometimes puts the figure label outside the download anchor.
            for anchor in soup.find_all("a", href=True):
                href = urljoin(response.url, anchor["href"])
                low = href.lower()
                if "xlsx" in low and ("7848" in low or "07848" in low):
                    parent_text = norm(anchor.parent.get_text(" ", strip=True))
                    if any(pattern.search(parent_text) for pattern in TARGET_PATTERNS):
                        candidates.append((parent_text, href))

        diagnostics.append(
            {
                "page": page_url,
                "status": response.status_code,
                "final_url": response.url,
                "candidate_count": len(candidates),
                "candidates": candidates,
            }
        )
        if candidates:
            label, href = candidates[0]
            print(
                json.dumps(
                    {
                        "attachment_label": label,
                        "attachment_url": href,
                        "page_diagnostics": diagnostics,
                    },
                    indent=2,
                )
            )
            return label, href

    raise RuntimeError(
        "Could not locate Source Data Fig. 3 attachment. Diagnostics:\n"
        + json.dumps(diagnostics, indent=2)
    )


def cell_value(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def inspect_workbook(xlsx_path: Path, outdir: Path, source_url: str) -> dict:
    workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    sheets = []
    preview_lines = []

    for index, ws in enumerate(workbook.worksheets, start=1):
        safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", ws.title).strip("_")
        if not safe_name:
            safe_name = f"sheet_{index}"
        csv_path = outdir / f"{index:02d}_{safe_name}.csv"

        nonempty = 0
        rows_written = 0
        numeric_count = 0
        text_tokens = []
        preview = []

        with csv_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = [cell_value(v) for v in row]
                if any(v != "" for v in values):
                    nonempty += 1
                    rows_written += 1
                    for v in values:
                        if isinstance(v, (int, float)) and not isinstance(v, bool):
                            numeric_count += 1
                        elif isinstance(v, str) and v:
                            text_tokens.append(v.lower())
                    if len(preview) < 25:
                        preview.append(values)
                writer.writerow(values)

        blob = " ".join(text_tokens)
        fig3a_score = sum(
            token in blob
            for token in (
                "3a",
                "ihsi",
                "+ir",
                "-ir",
                "co2",
                "time",
                "seconds",
                "second",
            )
        )
        sheets.append(
            {
                "index": index,
                "title": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "nonempty_rows": nonempty,
                "numeric_cell_count": numeric_count,
                "fig3a_keyword_score": fig3a_score,
                "csv": csv_path.name,
                "preview": preview,
            }
        )
        preview_lines.append(f"\n## SHEET {index}: {ws.title}\n")
        for row in preview:
            preview_lines.append("\t".join(map(str, row)) + "\n")

    manifest = {
        "source_url": source_url,
        "xlsx_file": xlsx_path.name,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "candidate_fig3a_sheets": [
            s["title"]
            for s in sorted(
                sheets,
                key=lambda x: (
                    x["fig3a_keyword_score"],
                    x["numeric_cell_count"],
                    x["nonempty_rows"],
                ),
                reverse=True,
            )[:3]
        ],
    }
    (outdir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    (outdir / "preview.txt").write_text(
        "".join(preview_lines),
        encoding="utf-8",
    )
    return manifest


def main() -> None:
    outdir = Path(sys.argv[1] if len(sys.argv) > 1 else "phase_a_chandel_fig3")
    outdir.mkdir(parents=True, exist_ok=True)

    label, attachment_url = find_attachment()
    response = fetch(attachment_url)
    xlsx_path = outdir / "Source_Data_Fig_3.xlsx"
    xlsx_path.write_bytes(response.content)

    if not response.content.startswith(b"PK"):
        raise RuntimeError(
            "Downloaded attachment does not look like an XLSX zip archive: "
            f"content-type={response.headers.get('content-type')} "
            f"final-url={response.url}"
        )

    manifest = inspect_workbook(
        xlsx_path,
        outdir,
        response.url,
    )
    manifest["attachment_label"] = label
    (outdir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
