"""Analyze Chandel et al. 2024 Nature Source Data Figure 3a.

The workbook contains a 3000-frame, 10-FPS aggregate IHSI time series for
the -IR and +IR zones. The public Arduino LongPulse controller matches the
published design of two 30-s CO2 pulses over a 300-s recording:

    0-60 s    CO2 off
    60-90 s   CO2 on
    90-180 s  CO2 off
    180-210 s CO2 on
    210-300 s CO2 off

This script is descriptive. Figure 3a contains the across-trial mean time
series, not the 18 individual trial trajectories, so it does not estimate
individual completion-time CDFs or infer timepoint-level uncertainty.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
from statistics import mean, median

from openpyxl import load_workbook


EPOCHS = (
    ("pre_first_pulse", 0.0, 60.0),
    ("first_co2_pulse", 60.0, 90.0),
    ("post_first_pulse", 90.0, 180.0),
    ("second_co2_pulse", 180.0, 210.0),
    ("post_second_pulse", 210.0, 300.0),
)

FRAME_RATE_HZ = 10.0
ROLLING_WINDOW_FRAMES = 50


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _load_fig3a(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Figure 3a"]
    headers = [ws.cell(3, i).value for i in range(1, 4)]
    expected = ["Frame", "Mean HSI No IR Zone", "Mean HSI IR Zone"]
    if headers != expected:
        raise ValueError(f"unexpected Figure 3a headers: {headers!r}")

    rows = []
    for frame, no_ir, ir in ws.iter_rows(
        min_row=4, max_col=3, values_only=True
    ):
        if frame is None:
            continue
        frame = int(frame)
        rows.append(
            {
                "frame": frame,
                "time_s": (frame - 1) / FRAME_RATE_HZ,
                "no_ir": float(no_ir),
                "ir": float(ir),
            }
        )
    if len(rows) != 3000:
        raise ValueError(f"expected 3000 frames, observed {len(rows)}")
    return rows


def _trapezoid(values, times):
    return sum(
        0.5 * (values[i] + values[i + 1]) * (times[i + 1] - times[i])
        for i in range(len(values) - 1)
    )


def _summarize(values, times):
    max_value = max(values)
    peak_index = values.index(max_value)
    return {
        "mean": mean(values),
        "median": median(values),
        "max": max_value,
        "peak_time_s": times[peak_index],
        "auc_trapezoid": _trapezoid(values, times),
    }


def _centered_rolling(values, window):
    n = len(values)
    left = window // 2
    right = window - left
    out = [None] * n
    for i in range(n):
        lo = max(0, i - left)
        hi = min(n, i + right)
        if hi - lo >= window // 2:
            out[i] = mean(values[lo:hi])
    return out


VALUE_KERNEL_TAUS_S = (5.0, 10.0, 20.0, 30.0, 60.0, 120.0)
CUMULATIVE_CUTS_S = (5.0, 10.0, 20.0, 30.0, 45.0, 60.0, 75.0, 90.0)


def _postpulse_timing_metrics(rows, start, end):
    block = [r for r in rows if start <= r["time_s"] < end]
    rel_time = [r["time_s"] - start for r in block]
    delta = [r["ir"] - r["no_ir"] for r in block]
    dt = 1.0 / FRAME_RATE_HZ

    signed_mass = sum(delta) * dt
    positive = [max(x, 0.0) for x in delta]
    positive_mass = sum(positive) * dt
    negative_mass = sum(max(-x, 0.0) for x in delta) * dt

    positive_centroid = (
        sum(t * x for t, x in zip(rel_time, positive)) / sum(positive)
        if sum(positive) > 0
        else None
    )

    cumulative_fraction = {}
    for cut in CUMULATIVE_CUTS_S:
        partial = sum(
            x for t, x in zip(rel_time, delta)
            if t < cut
        ) * dt
        cumulative_fraction[str(cut)] = (
            partial / signed_mass
            if signed_mass != 0
            else None
        )

    exponential_weighted_mean_delta = {}
    for tau in VALUE_KERNEL_TAUS_S:
        weights = [math.exp(-t / tau) for t in rel_time]
        exponential_weighted_mean_delta[str(tau)] = (
            sum(x * w for x, w in zip(delta, weights)) / sum(weights)
        )

    return {
        "signed_advantage_auc": signed_mass,
        "positive_advantage_auc": positive_mass,
        "negative_advantage_auc": negative_mass,
        "positive_advantage_centroid_s_after_pulse": positive_centroid,
        "cumulative_signed_advantage_fraction_by_s": cumulative_fraction,
        "exponential_value_kernel_weighted_mean_delta": exponential_weighted_mean_delta,
    }


def analyze(path: Path):
    rows = _load_fig3a(path)
    epochs = {}

    for name, start, end in EPOCHS:
        block = [r for r in rows if start <= r["time_s"] < end]
        times = [r["time_s"] for r in block]
        no_ir = [r["no_ir"] for r in block]
        ir = [r["ir"] for r in block]
        no_summary = _summarize(no_ir, times)
        ir_summary = _summarize(ir, times)
        epochs[name] = {
            "start_s": start,
            "end_s": end,
            "frames": len(block),
            "no_ir": no_summary,
            "ir": ir_summary,
            "ir_minus_no_ir_mean": ir_summary["mean"] - no_summary["mean"],
            "ir_to_no_ir_mean_ratio": (
                ir_summary["mean"] / no_summary["mean"]
            ),
        }

    no_all = [r["no_ir"] for r in rows]
    ir_all = [r["ir"] for r in rows]
    no_roll = _centered_rolling(no_all, ROLLING_WINDOW_FRAMES)
    ir_roll = _centered_rolling(ir_all, ROLLING_WINDOW_FRAMES)

    smooth_peaks = {}
    for name, start, end in (
        ("post_first_pulse", 90.0, 180.0),
        ("post_second_pulse", 210.0, 300.0),
    ):
        idxs = [
            i for i, r in enumerate(rows)
            if start <= r["time_s"] < end
        ]
        smooth_peaks[name] = {}
        for label, series in (("no_ir", no_roll), ("ir", ir_roll)):
            candidates = [(series[i], i) for i in idxs if series[i] is not None]
            peak_value, peak_i = max(candidates)
            peak_time = rows[peak_i]["time_s"]
            smooth_peaks[name][label] = {
                "peak_5s_mean": peak_value,
                "peak_time_s": peak_time,
                "seconds_after_pulse_end": peak_time - start,
            }

    contrasts = {
        "post_to_during_mean_ratio_first": {
            label: (
                epochs["post_first_pulse"][label]["mean"]
                / epochs["first_co2_pulse"][label]["mean"]
            )
            for label in ("no_ir", "ir")
        },
        "post_to_during_mean_ratio_second": {
            label: (
                epochs["post_second_pulse"][label]["mean"]
                / epochs["second_co2_pulse"][label]["mean"]
            )
            for label in ("no_ir", "ir")
        },
        "second_to_first_post_mean_ratio": {
            label: (
                epochs["post_second_pulse"][label]["mean"]
                / epochs["post_first_pulse"][label]["mean"]
            )
            for label in ("no_ir", "ir")
        },
        "second_to_first_pulse_mean_ratio": {
            label: (
                epochs["second_co2_pulse"][label]["mean"]
                / epochs["first_co2_pulse"][label]["mean"]
            )
            for label in ("no_ir", "ir")
        },
    }

    return {
        "status": "DESCRIPTIVE_AGGREGATE_TIMECOURSE_ONLY",
        "source": {
            "paper": "Chandel et al. 2024 Nature 633:615-623",
            "doi": "10.1038/s41586-024-07848-5",
            "source_data_url": (
                "https://media.springernature.com/original/"
                "springer-static/esm/art%3A10.1038%2Fs41586-024-07848-5/"
                "MediaObjects/41586_2024_7848_MOESM6_ESM.xlsx"
            ),
            "workbook_sha256": _sha256(path),
            "sheet": "Figure 3a",
            "frame_rate_hz": FRAME_RATE_HZ,
            "frame_count": len(rows),
            "duration_s": len(rows) / FRAME_RATE_HZ,
            "published_replication": (
                "mean IHSI across 18 trials "
                "(6 biological x 3 technical replicates)"
            ),
            "pulse_schedule_provenance": (
                "public Arduino_CO2_Release/PushbottonCO2_LongPulse.ino; "
                "schedule matches the manuscript's two 30-s CO2 pulses"
            ),
        },
        "epochs": epochs,
        "smoothed_peaks": smooth_peaks,
        "contrasts": contrasts,
        "postpulse_value_kernel_sensitivity": {
            "post_first_pulse": _postpulse_timing_metrics(rows, 90.0, 180.0),
            "post_second_pulse": _postpulse_timing_metrics(rows, 210.0, 300.0),
            "interpretation": (
                "These are descriptive weightings of the aggregate IR-minus-no-IR "
                "IHSI trajectory. They are not completion-time CDFs or fitness estimates."
            ),
        },
        "interpretation": {
            "primary": (
                "The +IR effect is much larger after each CO2 pulse than "
                "during the first pulse, consistent with sustained rather "
                "than merely faster initial host-seeking activity."
            ),
            "claim_ceiling": (
                "The workbook supplies an aggregate mean IHSI trajectory. "
                "It does not identify individual decision/completion times, "
                "trial-level uncertainty, fitness, or the adaptive-gain "
                "routing comparator."
            ),
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--output-json", type=Path)
    parser.add_argument("--output-epoch-csv", type=Path)
    args = parser.parse_args()

    result = analyze(args.xlsx)
    rendered = json.dumps(result, indent=2, sort_keys=True)
    print(rendered)

    if args.output_json:
        args.output_json.write_text(rendered + "\n", encoding="utf-8")

    if args.output_epoch_csv:
        with args.output_epoch_csv.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(
                [
                    "epoch",
                    "start_s",
                    "end_s",
                    "no_ir_mean",
                    "ir_mean",
                    "ir_minus_no_ir_mean",
                    "ir_to_no_ir_mean_ratio",
                ]
            )
            for name, _, _ in EPOCHS:
                e = result["epochs"][name]
                writer.writerow(
                    [
                        name,
                        e["start_s"],
                        e["end_s"],
                        e["no_ir"]["mean"],
                        e["ir"]["mean"],
                        e["ir_minus_no_ir_mean"],
                        e["ir_to_no_ir_mean_ratio"],
                    ]
                )


if __name__ == "__main__":
    main()
