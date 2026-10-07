"""Fetch and inventory Uehara et al. 2026 Supplementary Data 1."""
from __future__ import annotations

import csv
import json
from pathlib import Path

import requests
from openpyxl import load_workbook

URL = (
    "https://media.springernature.com/original/springer-static/esm/"
    "art%3A10.1038%2Fs42003-026-09987-z/MediaObjects/"
    "42003_2026_9987_MOESM3_ESM.xlsx"
)


def main() -> None:
    outdir = Path("uehara_supplementary_data")
    outdir.mkdir(exist_ok=True)
    xlsx = outdir / "42003_2026_9987_MOESM3_ESM.xlsx"

    response = requests.get(
        URL,
        headers={"User-Agent": "Mozilla/5.0 adaptive-gain-data-audit/1.0"},
        timeout=60,
    )
    response.raise_for_status()
    xlsx.write_bytes(response.content)
    if not response.content.startswith(b"PK"):
        raise RuntimeError(
            "Downloaded content does not look like XLSX: "
            f"{response.headers.get('content-type')}"
        )

    wb = load_workbook(xlsx, data_only=True, read_only=True)
    sheets = []
    for index, ws in enumerate(wb.worksheets, start=1):
        csv_path = outdir / f"{index:02d}_{ws.title.replace('/', '_')}.csv"
        nonempty_rows = 0
        max_nonempty_cols = 0
        preview = []
        numeric_cells = 0
        text_cells = 0
        with csv_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            for row in ws.iter_rows(values_only=True):
                values = list(row)
                writer.writerow("" if v is None else v for v in values)
                nonempty = [v for v in values if v is not None]
                if nonempty:
                    nonempty_rows += 1
                    max_nonempty_cols = max(max_nonempty_cols, len(nonempty))
                    if len(preview) < 20:
                        preview.append(values)
                for value in nonempty:
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        numeric_cells += 1
                    else:
                        text_cells += 1
        sheets.append(
            {
                "index": index,
                "title": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "nonempty_rows": nonempty_rows,
                "max_nonempty_cols": max_nonempty_cols,
                "numeric_cells": numeric_cells,
                "text_cells": text_cells,
                "csv": csv_path.name,
                "preview": preview,
            }
        )

    manifest = {
        "source_url": URL,
        "byte_size": xlsx.stat().st_size,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "frame_level_plausibility": {
            "expected_rows_for_9min_at_60fps": 32400,
            "expected_rows_for_8min_analysis_at_60fps": 28800,
            "any_sheet_ge_28800_rows": any(s["max_row"] >= 28800 for s in sheets),
            "any_sheet_ge_100_individual_columns": any(
                s["max_column"] >= 100 for s in sheets
            ),
        },
    }
    (outdir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
