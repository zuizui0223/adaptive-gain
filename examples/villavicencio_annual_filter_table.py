"""Join frozen conventional-filter covariates to annual Villavicencio dyad rows.

This step is deliberately effect-blind: it only constructs the analysis table
and audits missingness. It does not fit a model and it does not infer
decision-equivalence classes.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import Counter
from pathlib import Path


PREDICTORS = (
    "phenological_overlap",
    "flower_abundance",
    "corolla_length",
    "corolla_aperture",
    "height_mean",
    "body_length",
    "proboscis_length",
    "proboscis_width",
    "body_width",
    "body_thickness",
)


def _clean_number(value):
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if not text or text.lower() in {"na", "nan", "none"}:
            return None
        value = text
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return out if math.isfinite(out) else None


def _matrix_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    columns = [str(value).strip() for value in rows[0][1:]]
    out = {}
    for row in rows[1:]:
        if row[0] is None:
            continue
        plant = str(row[0]).strip()
        for pollinator, value in zip(columns, row[1:]):
            out[(plant, pollinator)] = _clean_number(value)
    return out


def _keyed_table(ws):
    rows = list(ws.iter_rows(values_only=True))
    headers = [
        "" if value is None else str(value).strip()
        for value in rows[0]
    ]
    out = {}
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip()
        out[key] = {
            headers[index]: _clean_number(row[index])
            for index in range(1, min(len(headers), len(row)))
        }
    return out


def _read_dyads(path: Path):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _bool(value: str) -> bool:
    text = value.strip().lower()
    if text in {"true", "1", "yes", "y"}:
        return True
    if text in {"false", "0", "no", "n"}:
        return False
    raise ValueError(f"unrecognized boolean value: {value!r}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("dyads", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    args = parser.parse_args()

    from openpyxl import load_workbook

    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    required = {
        "phenological_overlap",
        "Flower_abundance",
        "plant_traits",
        "pollinator_traits",
    }
    missing = sorted(required - set(wb.sheetnames))
    if missing:
        raise ValueError(f"workbook missing conventional-filter sheets: {missing!r}")

    phenology = _matrix_sheet(wb["phenological_overlap"])
    flower_abundance = _keyed_table(wb["Flower_abundance"])
    plant_traits = _keyed_table(wb["plant_traits"])
    pollinator_traits = _keyed_table(wb["pollinator_traits"])

    source_rows = _read_dyads(args.dyads)
    output_rows = []
    missing_counts = Counter()
    risk_counts = Counter()
    event_counts = Counter()

    for row in source_rows:
        plant = row["plant"]
        pollinator = row["pollinator"]
        previous_link = _bool(row["previous_link"])
        current_link = _bool(row["current_link"])

        if previous_link:
            risk_set = "loss"
            outcome = int(not current_link)
        else:
            risk_set = "gain"
            outcome = int(current_link)

        covariates = {
            "phenological_overlap": phenology.get((plant, pollinator)),
            "flower_abundance": flower_abundance.get(plant, {}).get("flower_abtot"),
            "corolla_length": plant_traits.get(plant, {}).get("corolla_length"),
            "corolla_aperture": plant_traits.get(plant, {}).get("corolla_aperture"),
            "height_mean": plant_traits.get(plant, {}).get("height_mean"),
            "body_length": pollinator_traits.get(pollinator, {}).get("body_length"),
            "proboscis_length": pollinator_traits.get(pollinator, {}).get("proboscis_length"),
            "proboscis_width": pollinator_traits.get(pollinator, {}).get("proboscis_width"),
            "body_width": pollinator_traits.get(pollinator, {}).get("body_width"),
            "body_thickness": pollinator_traits.get(pollinator, {}).get("body_thickness"),
        }

        for name, value in covariates.items():
            if value is None:
                missing_counts[name] += 1

        complete_case = all(covariates[name] is not None for name in PREDICTORS)
        transition = f"{row['previous_period']}->{row['current_period']}"
        risk_counts[risk_set] += 1
        event_counts[risk_set] += outcome

        output_rows.append(
            {
                **row,
                "transition": transition,
                "risk_set": risk_set,
                "outcome": outcome,
                **{
                    name: "" if value is None else value
                    for name, value in covariates.items()
                },
                "complete_case": complete_case,
            }
        )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    fields = list(output_rows[0].keys()) if output_rows else [
        "previous_period",
        "current_period",
        "plant",
        "pollinator",
        "transition",
        "risk_set",
        "outcome",
        *PREDICTORS,
        "complete_case",
    ]
    with args.output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(output_rows)

    total = len(output_rows)
    complete = sum(str(row["complete_case"]).lower() == "true" for row in output_rows)
    summary = {
        "schema": "adaptive-gain-villavicencio-conventional-filter-table-v1",
        "row_count": total,
        "complete_case_count": complete,
        "complete_case_fraction": None if total == 0 else complete / total,
        "predictors": list(PREDICTORS),
        "missing_counts": {
            name: missing_counts[name]
            for name in PREDICTORS
        },
        "coverage": {
            name: None if total == 0 else (total - missing_counts[name]) / total
            for name in PREDICTORS
        },
        "risk_sets": {
            risk: {
                "n": risk_counts[risk],
                "events": event_counts[risk],
                "nonevents": risk_counts[risk] - event_counts[risk],
                "event_fraction": (
                    None
                    if risk_counts[risk] == 0
                    else event_counts[risk] / risk_counts[risk]
                ),
            }
            for risk in ("gain", "loss")
        },
        "effect_fitted": False,
        "decision_equivalence_inferred": False,
    }
    args.summary.write_text(
        json.dumps(summary, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
