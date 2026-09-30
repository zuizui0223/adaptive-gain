"""Run a detection-sensitive annual Villavicencio Stage-1 audit from the verified XLSX.

This is a fallback feasibility analysis, not the frozen 18-subseason primary
design. Species presence is intentionally defined only by positive observed
links because the annual workbook does not independently establish pollinator
presence when degree is zero. The output therefore cannot be interpreted as
true species turnover without an additional detection/presence layer.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from dataclasses import asdict
from pathlib import Path

from adaptive_gain.empirical_rewiring import (
    rewiring_estimability_audit,
    transition_dyad_rows,
    transition_rewiring_receipt,
)


YEARS = tuple(range(2006, 2012))
REQUIRED_CONTEXT_SHEETS = (
    "insect_names",
    "plant_names",
    "Flower_abundance",
    "phenological_overlap",
    "plant_traits",
    "pollinator_traits",
)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _read_matrix(ws) -> dict[tuple[str, str], float]:
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2 or len(rows[0]) < 2:
        raise ValueError(f"sheet {ws.title!r} is not a usable interaction matrix")

    pollinators = [_cell_text(value) for value in rows[0][1:]]
    if any(not code for code in pollinators):
        raise ValueError(f"sheet {ws.title!r} has blank pollinator headers")

    network: dict[tuple[str, str], float] = {}
    for row in rows[1:]:
        plant = _cell_text(row[0] if row else None)
        if not plant:
            continue
        values = list(row[1:])
        if len(values) < len(pollinators):
            values.extend([None] * (len(pollinators) - len(values)))
        for pollinator, raw in zip(pollinators, values):
            if raw in (None, ""):
                continue
            weight = float(raw)
            if weight < 0:
                raise ValueError(f"negative interaction weight in {ws.title!r}")
            if weight > 0:
                network[(plant, pollinator)] = weight
    return network


def _read_code_column(ws, preferred_names: tuple[str, ...]) -> set[str]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return set()
    headers = [_cell_text(value).lower() for value in rows[0]]
    index = None
    for name in preferred_names:
        if name.lower() in headers:
            index = headers.index(name.lower())
            break
    if index is None:
        index = 0
    out = set()
    for row in rows[1:]:
        if index >= len(row):
            continue
        value = _cell_text(row[index])
        if value:
            out.add(value)
    return out


def _receipt_payload(receipt) -> dict:
    out = asdict(receipt)
    out.update(
        {
            "shared_species_rewiring_count": receipt.shared_species_rewiring_count,
            "species_turnover_link_count": receipt.species_turnover_link_count,
            "exact_partition": receipt.exact_partition,
            "link_jaccard_dissimilarity": receipt.link_jaccard_dissimilarity,
            "shared_species_rewiring_fraction_of_union": (
                receipt.shared_species_rewiring_fraction_of_union
            ),
            "species_turnover_fraction_of_union": (
                receipt.species_turnover_fraction_of_union
            ),
            "rewiring_opportunity_rate": receipt.rewiring_opportunity_rate,
        }
    )
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--dyad-output", type=Path)
    args = parser.parse_args()

    from openpyxl import load_workbook

    raw = args.workbook.read_bytes()
    source_sha256 = hashlib.sha256(raw).hexdigest()
    wb = load_workbook(args.workbook, read_only=True, data_only=True)

    required_annual = tuple(f"interactions_{year}" for year in YEARS)
    missing = [
        name
        for name in (*required_annual, *REQUIRED_CONTEXT_SHEETS)
        if name not in wb.sheetnames
    ]
    if missing:
        raise ValueError(f"workbook missing required sheets: {missing!r}")

    networks = {
        year: _read_matrix(wb[f"interactions_{year}"])
        for year in YEARS
    }

    plant_codes = _read_code_column(
        wb["plant_names"],
        ("plant_code", "plant.code", "code"),
    )
    pollinator_codes = _read_code_column(
        wb["insect_names"],
        ("pol_code", "pollinator_code", "code"),
    )

    observed_plants = {plant for network in networks.values() for plant, _ in network}
    observed_pollinators = {
        pollinator for network in networks.values() for _, pollinator in network
    }

    transitions = []
    transition_rows = {}
    dyad_export = []

    for previous_year, current_year in zip(YEARS, YEARS[1:]):
        previous = networks[previous_year]
        current = networks[current_year]
        receipt = transition_rewiring_receipt(previous, current)
        label = f"{previous_year}->{current_year}"
        rows = transition_dyad_rows(previous, current)
        transition_rows[label] = rows
        transitions.append(
            {
                "previous_period": str(previous_year),
                "current_period": str(current_year),
                "role": "annual_detection_sensitive_fallback",
                "receipt": _receipt_payload(receipt),
            }
        )
        if args.dyad_output is not None:
            for row in rows:
                dyad_export.append(
                    {
                        "previous_period": previous_year,
                        "current_period": current_year,
                        **asdict(row),
                    }
                )

    estimability = rewiring_estimability_audit(transition_rows)

    inventory = []
    for name in wb.sheetnames:
        ws = wb[name]
        inventory.append(
            {
                "sheet": name,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
            }
        )

    result = {
        "schema": "adaptive-gain-villavicencio-annual-stage1-fallback-v1",
        "source_file": args.workbook.name,
        "source_sha256": source_sha256,
        "analysis_status": "detection_sensitive_feasibility_only",
        "primary_18_subseason_design_replaced": False,
        "decision_equivalence_inferred": False,
        "presence_basis": "observed_positive_links",
        "presence_warning": (
            "zero observed degree is not treated as verified species absence; "
            "turnover and rewiring partition is detection-sensitive"
        ),
        "years": list(YEARS),
        "transition_count": len(transitions),
        "join_audit": {
            "plant_name_table_count": len(plant_codes),
            "pollinator_name_table_count": len(pollinator_codes),
            "observed_plant_code_count": len(observed_plants),
            "observed_pollinator_code_count": len(observed_pollinators),
            "matched_observed_plant_count": len(observed_plants & plant_codes),
            "unmatched_observed_plant_count": len(observed_plants - plant_codes),
            "matched_observed_pollinator_count": len(
                observed_pollinators & pollinator_codes
            ),
            "unmatched_observed_pollinator_count": len(
                observed_pollinators - pollinator_codes
            ),
        },
        "context_sheets_present": {
            name: name in wb.sheetnames for name in REQUIRED_CONTEXT_SHEETS
        },
        "estimability": {
            **asdict(estimability),
            "changed_fraction": estimability.changed_fraction,
            "global_outcome_nondegenerate": estimability.global_outcome_nondegenerate,
            "has_within_transition_contrast": (
                estimability.has_within_transition_contrast
            ),
        },
        "transitions": transitions,
        "workbook_inventory": inventory,
        "claim_ceiling": (
            "This annual fallback can establish response estimability only. "
            "It cannot validate environmental routeability, decision-structural "
            "turnover, or true species turnover without independent presence "
            "and decision-structure information."
        ),
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(result, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )

    if args.dyad_output is not None:
        args.dyad_output.parent.mkdir(parents=True, exist_ok=True)
        fields = [
            "previous_period",
            "current_period",
            "plant",
            "pollinator",
            "previous_weight",
            "current_weight",
            "previous_link",
            "current_link",
            "changed",
            "direction",
            "permitted",
        ]
        with args.dyad_output.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(dyad_export)


if __name__ == "__main__":
    main()
