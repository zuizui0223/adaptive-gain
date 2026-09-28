"""Audit outcome-support leakage in the Villavicencio phenology surface.

The published workbook phenological-overlap matrix was built for analyses of
observed interactions. Before reusing it as a candidate-dyad predictor for
link gains, this audit asks whether non-zero phenology is available for dyads
that never interact in any annual network.

The audit also reports held-out performance of the already-frozen model with
phenological_overlap removed. It does not refit any model.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from pathlib import Path


SCOPES = {
    "strict_core_2008_2011": 3,
    "near_core_2007_2011": 4,
    "all_annual_2006_2011": 5,
}


def _read_csv(path: Path):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _number(value):
    if value is None:
        return 0.0
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return 0.0
        value = text
    try:
        out = float(value)
    except (TypeError, ValueError):
        return 0.0
    return out if math.isfinite(out) else 0.0


def _matrix_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"empty matrix sheet: {ws.title}")
    columns = [str(value).strip() for value in rows[0][1:]]
    out = {}
    for row in rows[1:]:
        if row[0] is None:
            continue
        plant = str(row[0]).strip()
        for pollinator, value in zip(columns, row[1:]):
            out[(plant, pollinator)] = _number(value)
    return out


def _log_loss(y, p):
    total = 0.0
    for outcome, probability in zip(y, p):
        probability = min(max(float(probability), 1e-8), 1 - 1e-8)
        total -= outcome * math.log(probability) + (1 - outcome) * math.log(
            1 - probability
        )
    return total / len(y)


def _brier(y, p):
    return sum(
        (float(probability) - outcome) ** 2
        for outcome, probability in zip(y, p)
    ) / len(y)


def _auc(y, p):
    positives = sum(y)
    negatives = len(y) - positives
    if positives == 0 or negatives == 0:
        return None
    ordered = sorted(zip(p, y), key=lambda item: item[0])
    rank_sum = 0.0
    index = 0
    rank = 1
    while index < len(ordered):
        end = index + 1
        while end < len(ordered) and ordered[end][0] == ordered[index][0]:
            end += 1
        average_rank = (rank + (rank + end - index - 1)) / 2
        rank_sum += average_rank * sum(
            outcome for _, outcome in ordered[index:end]
        )
        rank += end - index
        index = end
    return (
        rank_sum - positives * (positives + 1) / 2
    ) / (positives * negatives)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("baseline_predictions", type=Path)
    parser.add_argument("ablation_predictions", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    from openpyxl import load_workbook

    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    phenology = _matrix_sheet(wb["phenological_overlap"])
    annual = [
        _matrix_sheet(wb[f"interactions_{year}"])
        for year in range(2006, 2012)
    ]
    key_sets = [set(phenology), *(set(matrix) for matrix in annual)]
    if any(keys != key_sets[0] for keys in key_sets[1:]):
        raise SystemExit(
            "phenology and annual matrices do not share one dyad surface"
        )

    keys = sorted(phenology)
    never_zero = never_positive = ever_zero = ever_positive = 0
    for dyad in keys:
        ever_linked = any(matrix[dyad] > 0 for matrix in annual)
        phen_positive = phenology[dyad] > 0
        if ever_linked and phen_positive:
            ever_positive += 1
        elif ever_linked:
            ever_zero += 1
        elif phen_positive:
            never_positive += 1
        else:
            never_zero += 1

    baseline = _read_csv(args.baseline_predictions)
    ablation = _read_csv(args.ablation_predictions)
    baseline_groups = defaultdict(list)
    for row in baseline:
        baseline_groups[
            (row["analysis_scope"], row["risk_set"], row["transition"])
        ].append(row)
    drop_groups = defaultdict(list)
    for row in ablation:
        if row["model"] == "drop_phenology":
            drop_groups[
                (row["analysis_scope"], row["risk_set"], row["transition"])
            ].append(row)

    performance = {}
    for scope, expected_fold_count in SCOPES.items():
        performance[scope] = {}
        for risk in ("gain", "loss"):
            transitions = sorted(
                transition
                for s, r, transition in baseline_groups
                if s == scope and r == risk
            )
            if len(transitions) != expected_fold_count:
                raise SystemExit(f"{scope}/{risk}: unexpected fold count")

            pooled_base = []
            pooled_drop = []
            folds = {}
            for transition in transitions:
                key = (scope, risk, transition)
                base_rows = baseline_groups[key]
                drop_rows = drop_groups[key]
                if len(base_rows) != len(drop_rows):
                    raise SystemExit(f"{key}: row-count mismatch")
                if [row["outcome"] for row in base_rows] != [
                    row["outcome"] for row in drop_rows
                ]:
                    raise SystemExit(f"{key}: outcome-order mismatch")
                pooled_base.extend(base_rows)
                pooled_drop.extend(drop_rows)

                y = [int(row["outcome"]) for row in base_rows]
                null_p = [
                    float(row["null_probability"]) for row in base_rows
                ]
                drop_p = [
                    float(row["probability"]) for row in drop_rows
                ]
                null_ll = _log_loss(y, null_p)
                drop_ll = _log_loss(y, drop_p)
                null_brier = _brier(y, null_p)
                drop_brier = _brier(y, drop_p)
                folds[transition] = {
                    "n": len(y),
                    "events": sum(y),
                    "drop_phenology_roc_auc": _auc(y, drop_p),
                    "drop_phenology_relative_log_loss_reduction_vs_null": (
                        (null_ll - drop_ll) / null_ll if null_ll else None
                    ),
                    "drop_phenology_brier_improvement_vs_null": (
                        null_brier - drop_brier
                    ),
                }

            y = [int(row["outcome"]) for row in pooled_base]
            null_p = [
                float(row["null_probability"]) for row in pooled_base
            ]
            drop_p = [
                float(row["probability"]) for row in pooled_drop
            ]
            null_ll = _log_loss(y, null_p)
            drop_ll = _log_loss(y, drop_p)
            null_brier = _brier(y, null_p)
            drop_brier = _brier(y, drop_p)
            performance[scope][risk] = {
                "n": len(y),
                "events": sum(y),
                "drop_phenology_log_loss": drop_ll,
                "null_log_loss": null_ll,
                "drop_phenology_relative_log_loss_reduction_vs_null": (
                    (null_ll - drop_ll) / null_ll if null_ll else None
                ),
                "drop_phenology_brier": drop_brier,
                "null_brier": null_brier,
                "drop_phenology_brier_improvement_vs_null": (
                    null_brier - drop_brier
                ),
                "drop_phenology_roc_auc": _auc(y, drop_p),
                "folds": folds,
            }

    total = len(keys)
    ever = ever_zero + ever_positive
    phen_positive_total = never_positive + ever_positive
    result = {
        "schema": (
            "adaptive-gain-villavicencio-phenology-support-leakage-audit-v1"
        ),
        "date": "2026-09-28",
        "status": "gain_filter_quarantined_support_leakage",
        "workbook_dyad_surface": {
            "dyad_count": total,
            "ever_linked_count": ever,
            "never_linked_count": never_zero + never_positive,
            "phenology_positive_count": phen_positive_total,
            "contingency": {
                "never_linked_phenology_zero": never_zero,
                "never_linked_phenology_positive": never_positive,
                "ever_linked_phenology_zero": ever_zero,
                "ever_linked_phenology_positive": ever_positive,
            },
            "probability_ever_linked_given_phenology_positive": (
                None
                if phen_positive_total == 0
                else ever_positive / phen_positive_total
            ),
            "probability_phenology_positive_given_ever_linked": (
                None if ever == 0 else ever_positive / ever
            ),
            "nonzero_phenology_subset_of_ever_linked_support": (
                never_positive == 0
            ),
        },
        "drop_phenology_performance": performance,
        "interpretation": {
            "support_audit": (
                "Every dyad with non-zero published phenological overlap occurs "
                "as a positive interaction in at least one annual network. No "
                "never-observed dyad has positive published overlap. The "
                "published overlap matrix therefore cannot be treated as an "
                "outcome-independent opportunity surface for candidate link "
                "gains in this reanalysis."
            ),
            "primary_gain": (
                "After removing phenological_overlap from the already-frozen "
                "model, primary gain discrimination falls to approximately "
                "chance and proper-score performance is slightly worse than "
                "the null. The earlier strong gain result is therefore "
                "quarantined rather than interpreted as ecological filter "
                "predictability."
            ),
            "primary_loss": (
                "The drop-phenology loss model retains only modest "
                "discrimination and proper-score improvement. Because all "
                "loss-risk dyads are already observed links and the remaining "
                "covariates include study-wide quantities, this remains "
                "descriptive rather than prospective or causal."
            ),
        },
        "not_a_claim_about_source_paper": (
            "This audit does not invalidate Peralta et al. (2020), which "
            "analyzed stability/functionality of observed interactions. It "
            "invalidates the present reuse of the published phenology surface "
            "as an independent predictor over the full candidate-dyad gain "
            "risk set."
        ),
        "required_next_step": (
            "Construct an opportunity/phenology surface independently from raw "
            "dated visitor and flower-observation records, excluding the focal "
            "dyad response, before making a gain-filter claim."
        ),
        "claim_ceiling": (
            "The annual full-model gain skill and zero-versus-positive "
            "phenology gate must not be used as evidence for ecological "
            "filtering, prospective forecasting, causal primacy, or routeability."
        ),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(result, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
